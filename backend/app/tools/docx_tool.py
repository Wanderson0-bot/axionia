from pathlib import Path
from docx import Document

def inspect_docx(path: str) -> dict:
    document = Document(path)
    return {"type":"DOCX", "size_bytes":Path(path).stat().st_size, "paragraphs":len(document.paragraphs), "tables":len(document.tables)}
