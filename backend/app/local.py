import json
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4
from app.db.connection import connection, open_pool, close_pool
from psycopg.types.json import Jsonb
from app.db.repositories import record_audit, add_message, replace_steps, update_steps_status, add_approval, add_invitation
import logging
import secrets, hashlib, hmac
from uuid import UUID
import pandas as pd
import fitz
from docx import Document
from fastapi import APIRouter, File, HTTPException, UploadFile, Header, Depends
from pydantic import BaseModel, Field, field_validator

ROOT=Path(__file__).resolve().parents[1]
UPLOADS=ROOT/'uploads'; OUTPUTS=ROOT/'outputs'
for d in (UPLOADS,OUTPUTS): d.mkdir(exist_ok=True)
ALLOWED={'.csv','.xlsx','.xls','.pdf','.docx','.png','.jpg','.jpeg'}; LIMIT=20*1024*1024
logger=logging.getLogger('axion.permissions')
PERMISSIONS=['view_workspace','upload_files','create_execution','approve_execution','view_students','view_grades','view_attendance','create_reports','export_personal_data','export_shared_data','manage_members','manage_permissions','view_institutional_dashboard','view_executions','create_executions','view_reports','manage_documents','view_indicators']
def stamp(): return datetime.now(timezone.utc)
def conn(): return connection()
def init():
 open_pool()
 with conn() as d:
  d.execute('SELECT 1 FROM users LIMIT 0')
  
def uuid_value(value):
 try:return UUID(str(value))
 except (ValueError,TypeError,AttributeError):return None
def permitted(user,space,permission):
 required={'create_execution':'create_executions','approve_execution':'approve_executions','view_institutional_dashboard':'view_executions'}.get(permission,permission)
 with conn() as d:
  r=d.execute("SELECT p.code FROM memberships m JOIN membership_permissions mp ON mp.membership_id=m.id JOIN permissions p ON p.id=mp.permission_id WHERE m.user_id=%s AND m.institution_id=%s AND COALESCE(m.status,'active')='active' AND p.legacy_only=FALSE",(uuid_value(user),uuid_value(space))).fetchall()
  member=d.execute("SELECT 1 FROM memberships WHERE user_id=%s AND institution_id=%s AND COALESCE(status,'active')='active'",(uuid_value(user),uuid_value(space))).fetchone()
 return bool(member and (permission=='view_workspace' or required in {x['code'] for x in r}))
def _is_institution_admin(user,space):
 with conn() as d:
  row=d.execute("SELECT 1 FROM memberships WHERE user_id=%s AND institution_id=%s AND COALESCE(status,'active')='active' AND role='admin'",(uuid_value(user),uuid_value(space))).fetchone()
 return bool(row)
def get_file(fid,space=None):
 with conn() as d:r=d.execute('SELECT * FROM files WHERE id=%s'+(' AND institution_id=%s' if space else ''),(uuid_value(fid),uuid_value(space)) if space else (uuid_value(fid),)).fetchone()
 if not r:return None
 x=dict(r);x.update(id=str(x['id']),workspace_id=str(x['institution_id']) if x['institution_id'] else None,user_id=str(x['uploaded_by']) if x['uploaded_by'] else None,extension=Path(x['name']).suffix.lower(),size=x['size_bytes'],stored_name=x['storage_path']);x['path']=str(Path(x['storage_path']) if Path(x['storage_path']).is_absolute() else (ROOT/ x['storage_path']));return x
def public_file(x):return {'id':str(x['id']),'name':x['name'],'type':x['extension'][1:],'size':x['size'],'status':'uploaded','created_at':x['created_at'],'workspace_id':x['workspace_id']}
def classify(command,files):
 t=command.lower(); exts={x['extension'] for x in files}; intent='unknown'; tool=''; needs=False; write=False
 if any(x in t for x in ['crie uma planilha','criar planilha','gere uma planilha']):intent,tool,write='create_spreadsheet','spreadsheet_tool',True
 elif 'relatório' in t or 'relatorio' in t:intent,tool,write='create_report','report_tool',True
 elif 'pdf' in t or '.pdf' in exts:intent,tool,needs='read_pdf','pdf_tool',True
 elif 'docx' in t or 'documento' in t or '.docx' in exts:intent,tool,needs='read_docx','docx_tool',True
 elif 'csv' in t or '.csv' in exts:intent,tool,needs='analyze_csv','csv_tool',True
 elif any(x in t for x in ['planilha','notas','excel','xlsx']) or exts&{'.xlsx','.xls'}:intent,tool,needs='analyze_spreadsheet','spreadsheet_tool',True
 elif any(x in t for x in ['analise','analisar','resuma','resumir']):intent,tool,needs='analyze_file','file_tool',True
 elif 'prova' in t:intent,write='create_exam',True
 elif 'simulado' in t:intent,write='create_simulado',True
 elif 'atividade' in t:intent,write='create_activity',True
 elif 'plano de estudo' in t:intent,write='study_plan',True
 qs=[]
 if intent=='unknown':qs=['Não identifiquei o tipo de tarefa. Você pode explicar o resultado esperado?']
 elif needs and not files:qs=['Envie o arquivo que devo utilizar para esta análise.']
 elif intent=='create_spreadsheet':qs=['Qual é o objetivo da planilha?','Quais informações ela deve conter?','Você deseja o formato XLSX ou CSV?']
 plan=[] if qs else ['Validar os arquivos selecionados.','Processar os dados com ferramenta local.','Validar e apresentar resultado estruturado.']
 if write and plan:plan.insert(2,'Criar um novo arquivo sem alterar o original.')
 return intent,tool,qs,plan,write
def dataframe_result(df):
 numeric=df.select_dtypes(include='number'); missing={str(k):int(v) for k,v in df.isna().sum().items() if v}
 means={str(k):round(float(v),2) for k,v in numeric.mean().items()}
 duplicates=int(df.duplicated().sum())
 insights=[f'Médias calculadas para {len(means)} coluna(s) numérica(s).',f'{duplicates} registro(s) duplicado(s) identificado(s).']+([f'{sum(missing.values())} valor(es) ausente(s) identificado(s).'] if missing else [])
 return {'summary':f'A planilha possui {len(df)} linhas e {len(df.columns)} colunas.','metrics':{'rows':int(len(df)),'columns':int(len(df.columns)),'missing_values':sum(missing.values()),'duplicate_records':duplicates},'insights':insights,'files':[],'warnings':[],'sources':[],'details':{'columns':[str(x) for x in df.columns],'missing_values':missing,'numeric_means':means,'duplicate_records':duplicates,'preview':df.head(10).fillna('').to_dict(orient='records')}}
def process(e,files):
 intent=e['intent']; results=[]
 for f in files:
  p=f['path']; ext=f['extension']
  if ext=='.csv':
   try: results.append(dataframe_result(pd.read_csv(p)))
   except UnicodeDecodeError:results.append(dataframe_result(pd.read_csv(p,encoding='latin1')))
  elif ext in {'.xlsx','.xls'}:
   sheets=pd.read_excel(p,sheet_name=None); first=next(iter(sheets.values())); r=dataframe_result(first);r['details']['sheets']=list(sheets.keys());results.append(r)
  elif ext=='.pdf':
   doc=fitz.open(p); text=''.join(page.get_text() for page in doc);results.append({'summary':f'PDF com {len(doc)} página(s).','metrics':{'pages':len(doc),'characters':len(text)},'insights':['Texto extraído com sucesso.' if text.strip() else 'O PDF não possui texto extraível.'],'files':[],'warnings':[],'sources':[],'details':{'metadata':doc.metadata or {},'text_preview':text[:3000]}});doc.close()
  elif ext=='.docx':
   doc=Document(p); paragraphs=[x.text for x in doc.paragraphs if x.text.strip()]; results.append({'summary':f'DOCX com {len(paragraphs)} parágrafo(s).','metrics':{'paragraphs':len(paragraphs),'words':sum(len(x.split()) for x in paragraphs)},'insights':['Leitura de parágrafos concluída.'],'files':[],'warnings':[],'sources':[],'details':{'titles':[x.text for x in doc.paragraphs if x.style and x.style.name.startswith('Heading')],'text_preview':'\n'.join(paragraphs)[:3000]}})
  elif ext in {'.png','.jpg','.jpeg'}:results.append({'summary':'Imagem registrada.','metrics':{},'insights':['Processamento visual avançado ainda não está habilitado.'],'files':[],'warnings':[],'sources':[]})
 if not results:return {'summary':'Nenhum arquivo processável foi associado a esta execução.','metrics':{},'insights':[],'files':[],'warnings':['Envie um arquivo compatível para obter análise.'],'sources':[]}
 if len(results)==1:return results[0]
 return {'summary':f'{len(results)} arquivos foram processados localmente.','metrics':{'files':len(results)},'insights':[],'files':results,'warnings':[],'sources':[]}
def public_exec(r):
 x=dict(r)
 for k in ['plan','questions','required_inputs','used_files','used_tools']:
  if isinstance(x.get(k),str):x[k]=json.loads(x[k])
  elif x.get(k) is None:x[k]=[]
 x['id']=str(x['id']);x['workspace_id']=str(x['workspace_id']) if x.get('workspace_id') else None;x['user_id']=str(x['user_id'])
 x['approval_required']=bool(x['approval_required']);return x
def authorized_workspace(user_id,workspace_id,permission='view_workspace'):
 if not permitted(user_id,workspace_id,permission): raise HTTPException(404,'Recurso não encontrado.')
def execution_for_user(execution_id,user_id):
 x=get_exec(execution_id)
 if not x or not permitted(user_id,x['workspace_id'],'view_workspace'): raise HTTPException(404,'Execução não encontrada.')
 return x
def get_exec(eid):
 with conn() as d:
  r=d.execute("SELECT e.id,e.user_id,e.institution_id AS workspace_id,e.objective AS command,e.status,e.intent,e.plan,e.questions,e.required_inputs,e.used_files,e.used_tools,e.approval_required,e.approved_at,e.created_at,e.updated_at,er.result_data AS result FROM executions e LEFT JOIN execution_results er ON er.execution_id=e.id WHERE e.id=%s",(uuid_value(eid),)).fetchone()
 return public_exec(r) if r else None
def update(eid,**values):
 values['updated_at']=stamp()
 result=values.pop('result',None) if 'result' in values else None
 columns={'command':'objective','workspace_id':'institution_id'}
 json_fields={'plan','questions','required_inputs','used_files','used_tools'}
 assignments=[];params=[]
 for key,value in values.items():
  column=columns.get(key,key)
  assignments.append(f'{column}=%s');params.append(Jsonb(value) if key in json_fields and value is not None else uuid_value(value) if key=='workspace_id' else value)
 with conn() as d:
  if assignments:d.execute('UPDATE executions SET '+','.join(assignments)+' WHERE id=%s',(*params,uuid_value(eid)))
  if result is not None:
   d.execute("INSERT INTO execution_results(id,execution_id,summary,result_data,created_at) VALUES(%s,%s,%s,%s,%s) ON CONFLICT(execution_id) DO UPDATE SET summary=EXCLUDED.summary,result_data=EXCLUDED.result_data,created_at=EXCLUDED.created_at",(uuid4(),uuid_value(eid),result.get('summary'),Jsonb(result),stamp()))
 return get_exec(eid)
class ExecutionRequest(BaseModel):command:str=Field(min_length=1,max_length=5000);workspace_id:str;file_ids:list[str]=Field(default_factory=list)
class Clarification(BaseModel):answers:dict[str,str]=Field(default_factory=dict);command:str|None=None;file_ids:list[str]=Field(default_factory=list)
class MemberInvite(BaseModel):name:str=Field(min_length=1);email:str=Field(min_length=3);role:str='member'
class MemberUpdate(BaseModel):role:str|None=None;status:str|None=None
class PermissionUpdate(BaseModel):
 permissions:list[str]=Field(default_factory=list)
 @field_validator('permissions')
 @classmethod
 def valid_permissions(cls,value):
  if not all(isinstance(item,str) and item.strip() for item in value): raise ValueError('permissions deve conter somente strings não vazias.')
  return list(dict.fromkeys(value))
router=APIRouter(prefix='/api')
def hash_password(password:str):
 salt=secrets.token_hex(16); digest=hashlib.pbkdf2_hmac('sha256',password.encode(),salt.encode(),120000).hex(); return salt+'$'+digest
def verify_password(password:str,stored:str):
 try:
  salt,digest=stored.split('$',1); return hmac.compare_digest(hashlib.pbkdf2_hmac('sha256',password.encode(),salt.encode(),120000).hex(),digest)
 except ValueError:return False
def authenticated_user(authorization:str|None=Header(default=None)):
 if not authorization or not authorization.startswith('Bearer '): raise HTTPException(401,'Sessão ausente ou inválida.')
 token=authorization[7:]
 token_hash=hashlib.sha256(token.encode()).hexdigest()
 with conn() as d: row=d.execute('SELECT user_id FROM sessions WHERE token_hash=%s AND revoked_at IS NULL AND expires_at>now()',(token_hash,)).fetchone()
 if not row: raise HTTPException(401,'Sessão expirada ou inválida.')
 return str(row['user_id'])
class Register(BaseModel):name:str=Field(min_length=1);email:str=Field(min_length=3);password:str=Field(min_length=8);account_type:str='individual';institution_name:str|None=None
class Login(BaseModel):email:str;password:str
@router.post('/auth/register')
def register(payload:Register):
 email=payload.email.lower().strip(); user_id=uuid4(); now=stamp(); current_workspace_id=None
 if payload.account_type not in ('individual','institution'): raise HTTPException(422,'Tipo de conta inválido.')
 with conn() as d:
  if d.execute('SELECT 1 FROM users WHERE email=%s',(email,)).fetchone(): raise HTTPException(409,'Este e-mail já está cadastrado.')
  d.execute("INSERT INTO users(id,name,email,password_hash,status,created_at) VALUES(%s,%s,%s,%s,'active',%s)",(user_id,payload.name.strip(),email,hash_password(payload.password),now))
  if payload.account_type == 'institution':
   current_workspace_id=uuid4(); workspace_name=(payload.institution_name or 'Instituição')
   d.execute("INSERT INTO institutions(id,name,type,created_by,created_at,updated_at) VALUES(%s,%s,'institution',%s,%s,%s)",(current_workspace_id,workspace_name,user_id,now,now))
   membership_id=uuid4();d.execute("INSERT INTO memberships(id,user_id,institution_id,role,status,joined_at) VALUES(%s,%s,%s,'admin','active',%s)",(membership_id,user_id,current_workspace_id,now))
   d.execute("INSERT INTO membership_permissions(id,membership_id,permission_id,granted_at) SELECT gen_random_uuid(),%s,id,%s FROM permissions WHERE legacy_only=FALSE",(membership_id,now))
   plan_id=d.execute("SELECT id FROM plans WHERE name='free' AND active=TRUE").fetchone()
   if plan_id:d.execute("INSERT INTO subscriptions(id,institution_id,plan_id,status,started_at) VALUES(%s,%s,%s,'active',%s)",(uuid4(),current_workspace_id,plan_id['id'],now))
  token=secrets.token_urlsafe(32); d.execute('INSERT INTO sessions(id,user_id,token_hash,expires_at,created_at) VALUES(%s,%s,%s,%s,%s)',(uuid4(),user_id,hashlib.sha256(token.encode()).hexdigest(),now+__import__('datetime').timedelta(days=30),now))
 return {'token':token,'user':{'id':str(user_id),'name':payload.name.strip(),'email':email,'current_workspace_id':str(current_workspace_id) if current_workspace_id else None,'account_type':payload.account_type,'plan':'free'}}
@router.post('/auth/login')
def login(payload:Login):
 with conn() as d:user=d.execute('SELECT * FROM users WHERE email=%s',(payload.email.lower().strip(),)).fetchone()
 if not user or not user['password_hash'] or not verify_password(payload.password,user['password_hash']): raise HTTPException(401,'E-mail ou senha inválidos.')
 token=secrets.token_urlsafe(32)
 now=stamp()
 with conn() as d:d.execute('INSERT INTO sessions(id,user_id,token_hash,expires_at,created_at) VALUES(%s,%s,%s,%s,%s)',(uuid4(),user['id'],hashlib.sha256(token.encode()).hexdigest(),now+__import__('datetime').timedelta(days=30),now))
 return {'token':token}
@router.post('/auth/logout')
def logout(authorization:str|None=Header(default=None)):
 if authorization and authorization.startswith('Bearer '):
  with conn() as d:d.execute('UPDATE sessions SET revoked_at=%s WHERE token_hash=%s AND revoked_at IS NULL',(stamp(),hashlib.sha256(authorization[7:].encode()).hexdigest()))
 return {'status':'logged_out'}
@router.post('/files/upload')
async def upload(files:list[UploadFile]=File(...),workspace_id:str='',user_id:str=Depends(authenticated_user)):
 if not workspace_id: raise HTTPException(422,'workspace_id é obrigatório.')
 if not permitted(user_id,workspace_id,'upload_files'):raise HTTPException(403,'Sem permissão para enviar arquivos neste workspace.')
 saved=[]
 for f in files:
  name=Path(f.filename or 'arquivo').name;ext=Path(name).suffix.lower();content=await f.read()
  if ext not in ALLOWED:raise HTTPException(415,f'Formato não aceito: {name}')
  if not content:raise HTTPException(422,f'Arquivo vazio: {name}')
  if len(content)>LIMIT:raise HTTPException(413,'Arquivo excede o limite de 20 MB.')
  fid=uuid4();stored=uuid4().hex+ext
  (UPLOADS/stored).write_bytes(content)
  with conn() as d: d.execute('INSERT INTO files(id,institution_id,uploaded_by,name,mime_type,size_bytes,storage_path,status,created_at,updated_at) VALUES(%s,%s,%s,%s,%s,%s,%s,\'available\',%s,%s)',(fid,uuid_value(workspace_id),uuid_value(user_id),name,f.content_type,len(content),'uploads/'+stored,stamp(),stamp())); record_audit(d,stamp,'file.uploaded','file',fid,workspace_id,user_id)
  saved.append(public_file(get_file(fid)))
 return saved[0] if len(saved)==1 else {'files':saved}
@router.get('/files')
def files(workspace_id:str='',user_id:str=Depends(authenticated_user)):
 if not workspace_id: raise HTTPException(422,'workspace_id é obrigatório.')
 if not permitted(user_id,workspace_id,'view_workspace'):raise HTTPException(403,'Sem acesso ao workspace.')
 with conn() as d:rows=d.execute('SELECT * FROM files WHERE institution_id=%s ORDER BY created_at DESC',(uuid_value(workspace_id),)).fetchall()
 return [public_file(get_file(str(r['id']),workspace_id)) for r in rows]
@router.post('/executions')
def create(req:ExecutionRequest,user_id:str=Depends(authenticated_user)):
 authorized_workspace(user_id,req.workspace_id,'create_execution')
 fs=[]
 for fid in req.file_ids:
  f=get_file(fid,req.workspace_id)
  if not f:raise HTTPException(404,f'Arquivo não encontrado no workspace: {fid}')
  fs.append(f)
 intent,tool,qs,plan,approval=classify(req.command,fs); status='pending_clarification' if qs else 'awaiting_approval'
 eid=uuid4();t=stamp();used_files=[{'id':str(f['id']),'name':f['name']} for f in fs]
 with conn() as d:
  d.execute("INSERT INTO executions(id,institution_id,user_id,objective,status,created_at,updated_at,intent,plan,questions,required_inputs,used_files,used_tools,approval_required) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,TRUE)",(eid,uuid_value(req.workspace_id),uuid_value(user_id),req.command,status,t,t,intent,Jsonb(plan),Jsonb(qs),Jsonb([]),Jsonb(used_files),Jsonb([tool] if tool else []))); replace_steps(d,stamp,eid,plan); add_message(d,stamp,eid,'user',req.command,user_id); add_message(d,stamp,eid,'system','Execução criada.'); record_audit(d,stamp,'execution.created','execution',eid,req.workspace_id,user_id)
 return get_exec(str(eid))
@router.get('/executions')
def history(workspace_id:str|None=None,status:str|None=None,intent:str|None=None,user_id:str=Depends(authenticated_user)):
 if not workspace_id: raise HTTPException(422,"workspace_id é obrigatório.")
 authorized_workspace(user_id,workspace_id)
 q=[];v=[]
 if workspace_id:q.append('institution_id=%s');v.append(uuid_value(workspace_id))
 if status:q.append('status=%s');v.append(status)
 if intent:q.append('intent=%s');v.append(intent)
 with conn() as d:rows=d.execute("SELECT e.id,e.user_id,e.institution_id AS workspace_id,e.objective AS command,e.status,e.intent,e.plan,e.questions,e.required_inputs,e.used_files,e.used_tools,e.approval_required,e.approved_at,e.created_at,e.updated_at,er.result_data AS result FROM executions e LEFT JOIN execution_results er ON er.execution_id=e.id"+((' WHERE '+' AND '.join(q)) if q else '')+' ORDER BY e.created_at DESC',v).fetchall()
 return [public_exec(r) for r in rows]
@router.get('/executions/{eid}')
def execution(eid:str,user_id:str=Depends(authenticated_user)):
 x=execution_for_user(eid,user_id)
 return x
@router.post('/executions/{eid}/clarify')
def clarify(eid:str,payload:Clarification,user_id:str=Depends(authenticated_user)):
 x=execution_for_user(eid,user_id)
 if x['status']!='pending_clarification':raise HTTPException(409,'Esta execução não está aguardando esclarecimentos.')
 command=payload.command or x['command']+' '+ ' '.join(payload.answers.values())
 existing_ids=[f['id'] for f in x['used_files']]; all_ids=list(dict.fromkeys(existing_ids+payload.file_ids)); fs=[]
 for file_id in all_ids:
  file=get_file(file_id,x['workspace_id'])
  if not file: raise HTTPException(404,'Arquivo informado não foi encontrado neste workspace.')
  fs.append(file)
 intent,tool,qs,plan,_=classify(command,fs)
 with conn() as d: replace_steps(d,stamp,eid,plan); add_message(d,stamp,eid,'user',' '.join(payload.answers.values()),user_id); record_audit(d,stamp,'execution.clarified','execution',eid,x['workspace_id'],user_id)
 return update(eid,command=command,intent=intent,questions=qs,plan=plan,used_files=[{'id':str(f['id']),'name':f['name']} for f in fs],used_tools=[tool] if tool else [],status='pending_clarification' if qs else 'awaiting_approval')
@router.post('/executions/{eid}/approve')
def approve(eid:str,user_id:str=Depends(authenticated_user)):
 x=get_exec(eid)
 if not x:raise HTTPException(404,'Execução não encontrada.')
 if not permitted(user_id,x['workspace_id'],'approve_execution'):raise HTTPException(403,'Sem permissão para aprovar.')
 if x['status']!='awaiting_approval':raise HTTPException(409,'A execução não está aguardando aprovação.')
 with conn() as d: add_approval(d,stamp,eid,user_id,'approved'); record_audit(d,stamp,'execution.approved','execution',eid,x['workspace_id'],user_id)
 return update(eid,status='approved',approved_at=stamp())
@router.post('/executions/{eid}/cancel')
def cancel(eid:str,user_id:str=Depends(authenticated_user)):
 x=execution_for_user(eid,user_id)
 if not permitted(user_id,x['workspace_id'],'create_execution'):raise HTTPException(403,'Sem permissão para cancelar.')
 if x['status'] in ['completed','cancelled']:raise HTTPException(409,'Esta execução não pode ser cancelada.')
 with conn() as d: add_approval(d,stamp,eid,user_id,'cancelled'); update_steps_status(d,stamp,eid,'cancelled'); record_audit(d,stamp,'execution.cancelled','execution',eid,x['workspace_id'],user_id)
 return update(eid,status='cancelled')
@router.post('/executions/{eid}/run')
def run(eid:str,user_id:str=Depends(authenticated_user)):
 x=get_exec(eid)
 if not x:raise HTTPException(404,'Execução não encontrada.')
 if x['status']!='approved':raise HTTPException(409,'Aprovação humana é obrigatória antes da execução.')
 if not permitted(user_id,x['workspace_id'],'create_execution'):raise HTTPException(403,'Sem permissão para executar.')
 fs=[]
 for item in x['used_files']:
  f=get_file(item['id'],x['workspace_id'])
  if not f:raise HTTPException(404,'Um arquivo associado não existe mais.')
  fs.append(f)
 with conn() as d: update_steps_status(d,stamp,eid,'running'); record_audit(d,stamp,'execution.started','execution',eid,x['workspace_id'],user_id)
 update(eid,status='running')
 try:
  r=process(x,fs)
  with conn() as d: update_steps_status(d,stamp,eid,'completed'); add_message(d,stamp,eid,'assistant',r['summary']); record_audit(d,stamp,'execution.completed','execution',eid,x['workspace_id'],user_id)
  return update(eid,status='completed',result=r)
 except Exception as error:update(eid,status='failed',result={'summary':'Falha no processamento local.','metrics':{},'insights':[],'files':[],'warnings':[str(error)],'sources':[]});raise HTTPException(422,'Falha ao processar arquivo: '+str(error))
@router.get('/workspaces')
def workspaces(user_id:str=Depends(authenticated_user)):
 with conn() as d: rows=d.execute("SELECT i.id,i.name,'institution' AS type,i.created_at FROM institutions i JOIN memberships m ON m.institution_id=i.id WHERE m.user_id=%s AND COALESCE(m.status,'active')='active' ORDER BY i.created_at DESC",(uuid_value(user_id),)).fetchall()
 return [{'id':str(r['id']),'name':r['name'],'type':r['type'],'created_at':r['created_at']} for r in rows]

@router.get('/workspaces/{workspace_id}/permissions')
def permissions(workspace_id:str,user_id:str=Depends(authenticated_user)):
 with conn() as d:
  row=d.execute("SELECT m.id,m.role FROM memberships m WHERE m.user_id=%s AND m.institution_id=%s AND COALESCE(m.status,'active')='active'",(uuid_value(user_id),uuid_value(workspace_id))).fetchone()
  if not row: raise HTTPException(403,'Sem acesso ao workspace.')
  codes=d.execute("SELECT p.code FROM membership_permissions mp JOIN permissions p ON p.id=mp.permission_id WHERE mp.membership_id=%s AND p.legacy_only=FALSE ORDER BY p.code",(row['id'],)).fetchall()
 return {'workspace_id':workspace_id,'user_id':user_id,'permissions':[x['code'] for x in codes],'is_admin':row['role']=='admin'}

@router.get('/me')
def me(user_id:str=Depends(authenticated_user)):
 with conn() as d:
  user=d.execute('SELECT id,name,email,created_at FROM users WHERE id=%s',(uuid_value(user_id),)).fetchone()
  current=d.execute("SELECT m.institution_id FROM memberships m WHERE m.user_id=%s AND COALESCE(m.status,'active')='active' ORDER BY m.joined_at DESC NULLS LAST LIMIT 1",(uuid_value(user_id),)).fetchone()
 if not user: raise HTTPException(404,'Usuário não encontrado.')
 return {'id':str(user['id']),'name':user['name'] or None,'email':user['email'] or None,'account_type':'individual','created_at':user['created_at'],'current_workspace_id':str(current['institution_id']) if current else None}

def institutional(workspace_id:str,user_id:str,permission='view_workspace'):
 with conn() as d: workspace=d.execute('SELECT * FROM institutions WHERE id=%s',(uuid_value(workspace_id),)).fetchone()
 if not workspace: raise HTTPException(404,'Workspace não encontrado.')
 if not permitted(user_id,workspace_id,permission): raise HTTPException(403,'Sem permissão para este workspace.')
 return workspace

def member_public(row):
 permissions=row.get('permissions') or []
 return {'id':str(row['id']),'workspace_id':str(row['workspace_id']),'user_id':str(row['user_id']) if row.get('user_id') else None,'name':row.get('name'),'email':row['email'],'role':row['role'],'status':row['status'],'created_at':row['created_at'],'permissions':permissions}

@router.get('/workspaces/{workspace_id}/members')
def members(workspace_id:str,user_id:str=Depends(authenticated_user)):
 institutional(workspace_id,user_id)
 with conn() as d:
  active=d.execute("SELECT m.id,m.institution_id AS workspace_id,m.user_id,u.name,u.email,m.role,COALESCE(m.status,'active') AS status,m.joined_at AS created_at,COALESCE(array_agg(p.code) FILTER(WHERE p.code IS NOT NULL AND p.legacy_only=FALSE),'{}') AS permissions FROM memberships m JOIN users u ON u.id=m.user_id LEFT JOIN membership_permissions mp ON mp.membership_id=m.id LEFT JOIN permissions p ON p.id=mp.permission_id WHERE m.institution_id=%s AND COALESCE(m.status,'active')<>'removed' GROUP BY m.id,u.id",(uuid_value(workspace_id),)).fetchall()
  pending=d.execute("SELECT id,institution_id AS workspace_id,NULL::uuid AS user_id,split_part(invited_email,'@',1) AS name,invited_email AS email,role,status,created_at,ARRAY[]::text[] AS permissions FROM invitations WHERE institution_id=%s AND status='pending' ORDER BY created_at",(uuid_value(workspace_id),)).fetchall()
 return [member_public(r) for r in active+pending]

@router.post('/workspaces/{workspace_id}/members/invite')
def invite(workspace_id:str,payload:MemberInvite,user_id:str=Depends(authenticated_user)):
 institutional(workspace_id,user_id,'manage_members')
 if payload.role not in ('admin','member'): raise HTTPException(422,'Papel de membro inválido.')
 email=payload.email.lower().strip()
 with conn() as d:
  existing=d.execute("SELECT 1 FROM memberships m JOIN users u ON u.id=m.user_id WHERE m.institution_id=%s AND lower(u.email)=%s AND COALESCE(m.status,'active')='active' UNION ALL SELECT 1 FROM invitations WHERE institution_id=%s AND lower(invited_email)=%s AND status='pending' LIMIT 1",(uuid_value(workspace_id),email,uuid_value(workspace_id),email)).fetchone()
  if existing: raise HTTPException(409,'Já existe um membro ou convite com este e-mail neste workspace.')
  invitation=add_invitation(d,stamp,workspace_id,email,payload.role,user_id)
  record_audit(d,stamp,'member.invited','invitation',invitation['id'],workspace_id,user_id)
 return {'id':invitation['id'],'workspace_id':workspace_id,'user_id':None,'name':payload.name,'email':email,'role':payload.role,'status':'pending','created_at':invitation['created_at'],'permissions':[]}

@router.patch('/workspaces/{workspace_id}/members/{member_id}')
def update_member(workspace_id:str,member_id:str,payload:MemberUpdate,user_id:str=Depends(authenticated_user)):
 institutional(workspace_id,user_id,'manage_members')
 if payload.role and payload.role not in ('admin','member'): raise HTTPException(422,'Papel de membro inválido.')
 if payload.status and payload.status not in ('active','pending','removed'): raise HTTPException(422,'Status inválido.')
 with conn() as d:
  row=d.execute("SELECT id FROM memberships WHERE id=%s AND institution_id=%s",(uuid_value(member_id),uuid_value(workspace_id))).fetchone()
  if row:
   d.execute("UPDATE memberships SET role=COALESCE(%s,role),status=COALESCE(%s,status) WHERE id=%s AND institution_id=%s",(payload.role,payload.status,uuid_value(member_id),uuid_value(workspace_id)))
  else:
   invitation_status={'pending':'pending','removed':'cancelled','active':'accepted'}.get(payload.status) if payload.status else None
   changed=d.execute("UPDATE invitations SET role=COALESCE(%s,role),status=COALESCE(%s,status),updated_at=%s WHERE id=%s AND institution_id=%s AND status='pending'",(payload.role,invitation_status,stamp(),uuid_value(member_id),uuid_value(workspace_id))).rowcount
   if not changed: raise HTTPException(404,'Membro não encontrado neste workspace.')
 return {'id':member_id,'workspace_id':workspace_id,'role':payload.role,'status':payload.status or 'active'}

@router.delete('/workspaces/{workspace_id}/members/{member_id}')
def remove_member(workspace_id:str,member_id:str,user_id:str=Depends(authenticated_user)):
 institutional(workspace_id,user_id,'manage_members')
 with conn() as d:
  changed=d.execute("UPDATE memberships SET status='removed',left_at=%s WHERE id=%s AND institution_id=%s",(stamp(),uuid_value(member_id),uuid_value(workspace_id))).rowcount
  if not changed: changed=d.execute("UPDATE invitations SET status='cancelled',updated_at=%s WHERE id=%s AND institution_id=%s AND status='pending'",(stamp(),uuid_value(member_id),uuid_value(workspace_id))).rowcount
 if not changed: raise HTTPException(404,'Membro não encontrado neste workspace.')
 return {'status':'removed','id':member_id}

@router.put('/workspaces/{workspace_id}/members/{member_id}/permissions')
def update_permissions(workspace_id:str,member_id:str,payload:PermissionUpdate,user_id:str=Depends(authenticated_user)):
 institutional(workspace_id,user_id,'manage_permissions')
 with conn() as d:
  membership=d.execute("SELECT id FROM memberships WHERE id=%s AND institution_id=%s AND COALESCE(status,'active')='active'",(uuid_value(member_id),uuid_value(workspace_id))).fetchone()
  if not membership: raise HTTPException(404,'Membro não encontrado neste workspace.')
  allowed=d.execute("SELECT code,id FROM permissions WHERE legacy_only=FALSE AND code=ANY(%s)",(payload.permissions,)).fetchall()
  if {x['code'] for x in allowed}!=set(payload.permissions): raise HTTPException(422,'Permissões inexistentes ou legadas não podem ser concedidas.')
  d.execute('DELETE FROM membership_permissions WHERE membership_id=%s',(membership['id'],))
  for permission in allowed: d.execute('INSERT INTO membership_permissions(id,membership_id,permission_id,granted_at) VALUES(%s,%s,%s,%s)',(uuid4(),membership['id'],permission['id'],stamp()))
 return {'workspace_id':workspace_id,'member_id':member_id,'permissions':payload.permissions}

@router.get('/workspaces/{workspace_id}/indicators')
def indicators(workspace_id:str,user_id:str=Depends(authenticated_user)):
 institutional(workspace_id,user_id,'view_institutional_dashboard')
 with conn() as d:
  members_total=d.execute("SELECT COUNT(*) AS n FROM memberships WHERE institution_id=%s AND COALESCE(status,'active')<>'removed'",(uuid_value(workspace_id),)).fetchone()['n']
  members_active=d.execute("SELECT COUNT(*) AS n FROM memberships WHERE institution_id=%s AND COALESCE(status,'active')='active'",(uuid_value(workspace_id),)).fetchone()['n']
  total=d.execute('SELECT COUNT(*) AS n FROM executions WHERE institution_id=%s',(uuid_value(workspace_id),)).fetchone()['n']
  completed=d.execute("SELECT COUNT(*) AS n FROM executions WHERE institution_id=%s AND status='completed'",(uuid_value(workspace_id),)).fetchone()['n']
  running=d.execute("SELECT COUNT(*) AS n FROM executions WHERE institution_id=%s AND status='running'",(uuid_value(workspace_id),)).fetchone()['n']
  files_count=d.execute('SELECT COUNT(*) AS n FROM files WHERE institution_id=%s',(uuid_value(workspace_id),)).fetchone()['n']
  last=d.execute('SELECT MAX(updated_at) AS last FROM executions WHERE institution_id=%s',(uuid_value(workspace_id),)).fetchone()['last']
 return {'members_total':members_total,'members_active':members_active,'executions_total':total,'executions_completed':completed,'executions_running':running,'files_processed':files_count,'last_activity':last}

@router.get('/workspaces/{workspace_id}/plan')
def plan(workspace_id:str,user_id:str=Depends(authenticated_user)):
 institutional(workspace_id,user_id)
 with conn() as d:
  members_count=d.execute("SELECT COUNT(*) AS n FROM memberships WHERE institution_id=%s AND COALESCE(status,'active')<>'removed'",(uuid_value(workspace_id),)).fetchone()['n']
  executions_count=d.execute('SELECT COUNT(*) AS n FROM executions WHERE institution_id=%s',(uuid_value(workspace_id),)).fetchone()['n']
  storage=d.execute('SELECT COALESCE(SUM(size_bytes),0) AS n FROM files WHERE institution_id=%s',(uuid_value(workspace_id),)).fetchone()['n']
  row=d.execute("SELECT p.name AS plan,s.status,p.limits FROM subscriptions s JOIN plans p ON p.id=s.plan_id WHERE s.institution_id=%s ORDER BY s.started_at DESC LIMIT 1",(uuid_value(workspace_id),)).fetchone()
  limits=row['limits'] if row else {}
 return {'plan':row['plan'] if row else 'free','status':row['status'] if row else 'active','limits':limits or {},'usage':{'members':members_count,'executions':executions_count,'storage_mb':round(storage/1024/1024,2)}}
